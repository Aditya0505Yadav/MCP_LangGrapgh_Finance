"""Build the fixed-contract Word report (Times New Roman 12pt body / 14pt headings,
embedded price chart, dashboard hyperlink) from a run payload. Returns bytes.
"""

from __future__ import annotations

import io
from typing import Any

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BODY_FONT, BODY_PT, HEAD_PT, TITLE_PT = "Times New Roman", 12, 14, 16


def _usd(n): return "—" if n is None else f"${n:,.2f}"
def _pct(n): return "—" if n is None else f"{n:+.2f}%"
def _big(n):
    if n is None:
        return "—"
    a = abs(n)
    if a >= 1e12: return f"${n/1e12:.2f}T"
    if a >= 1e9: return f"${n/1e9:.2f}B"
    if a >= 1e6: return f"${n/1e6:.2f}M"
    return _usd(n)


def _chart_png(payload: dict) -> io.BytesIO:
    obs = payload["observations"]
    prices = [o.get("adjusted_close") or o.get("close") for o in obs]
    dates = [o["date"] for o in obs]
    fig, ax = plt.subplots(figsize=(9.5, 3.3), dpi=130)
    ax.plot(range(len(prices)), prices, color="#2a5298", linewidth=1.8)
    if prices:
        ax.fill_between(range(len(prices)), prices, min(prices), color="#2a5298", alpha=0.08)
    ax.set_title(f"{payload['symbol']} adjusted close — {dates[0]} to {dates[-1]}"
                 if dates else payload["symbol"], fontsize=11, color="#1a2436")
    step = max(1, len(dates) // 8)
    ax.set_xticks(range(0, len(dates), step))
    ax.set_xticklabels([dates[i] for i in range(0, len(dates), step)], rotation=30, fontsize=8)
    ax.tick_params(colors="#5d6d7e", labelsize=8)
    ax.grid(True, color="#e3ebf0", linewidth=0.7)
    for s in ax.spines.values():
        s.set_color("#e3ebf0")
    ax.set_ylabel("USD", fontsize=9, color="#5d6d7e")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def build_report_docx(payload: dict, dashboard_url: str | None = None,
                      generated_on: str = "") -> bytes:
    q, f = payload["quote"], payload["fundamentals"]
    sm, fm = payload["stock_metrics"], payload["financial_metrics"]
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    def rf(run, size=BODY_PT, bold=False, italic=False, color=None):
        run.font.name = BODY_FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        if color:
            run.font.color.rgb = color
        el = run._element
        el.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        el.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)

    def title(t):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rf(p.add_run(t), TITLE_PT, bold=True)

    def heading(t):
        p = doc.add_paragraph()
        rf(p.add_run(t), HEAD_PT, bold=True, color=RGBColor(0x1E, 0x3C, 0x72))

    def para(t, italic=False):
        rf(doc.add_paragraph().add_run(t), italic=italic)

    def bullet(t):
        rf(doc.add_paragraph(style="List Bullet").add_run(t))

    def hyperlink(paragraph, url, text):
        part = paragraph.part
        rid = part.relate_to(
            url, "http://schemas.openxmlformats.org/officeDocument/2006/"
                 "relationships/hyperlink", is_external=True)
        link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), rid)
        r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(qn("w:ascii"), BODY_FONT); rfonts.set(qn("w:hAnsi"), BODY_FONT); rpr.append(rfonts)
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(BODY_PT * 2)); rpr.append(sz)
        col = OxmlElement("w:color"); col.set(qn("w:val"), "0563C1"); rpr.append(col)
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
        r.append(rpr)
        t = OxmlElement("w:t"); t.text = text; r.append(t)
        link.append(r); paragraph._p.append(link)

    def kv(rows):
        t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
        for k, v in rows:
            cells = t.add_row().cells
            rf(cells[0].paragraphs[0].add_run(k), bold=True)
            rf(cells[1].paragraphs[0].add_run(v))
        doc.add_paragraph()

    title(f"{payload['company_name']} ({payload['symbol']}) — Real-Time Investment Analysis")
    meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf(meta.add_run(f"Generated {generated_on} · data as of {q.get('data_as_of')} · "
                    f"model {payload.get('model')} · mode {payload.get('data_mode')} · "
                    f"source: Yahoo Finance (live)"), size=10, color=RGBColor(0x5D, 0x6D, 0x7E))

    heading("1. Interactive Dashboard & Charts (AI-generated)")
    p = doc.add_paragraph(); rf(p.add_run("Live dashboard with interactive price chart and KPIs: "))
    if dashboard_url:
        hyperlink(p, dashboard_url, "Open the interactive dashboard")
    else:
        rf(p.add_run("(served by this app at /analyses/<run_id>/dashboard)"))

    heading("2. Real-Time Snapshot")
    kv([
        ("Current price", f"{_usd(q.get('current_price'))} ({q.get('currency')})"),
        ("Day change", f"{q['day_change']:+.2f} ({_pct(q.get('day_change_pct'))})"
                       if q.get("day_change") is not None else "—"),
        ("Market cap", _big(f.get("market_cap"))),
        ("P/E (trailing)", f"{f['pe_ratio']:.2f}" if f.get("pe_ratio") else "—"),
        ("52-week range", f"{_usd(q.get('fifty_two_week_low'))} – {_usd(q.get('fifty_two_week_high'))}"),
        ("As of", str(q.get("data_as_of"))),
    ])

    heading("3. Price Performance")
    doc.add_picture(_chart_png(payload), width=Inches(6.3))
    doc.add_paragraph()
    kv([
        ("Start / End", f"{_usd(sm.get('start_price'))}  →  {_usd(sm.get('end_price'))}"),
        ("Period return", _pct(sm.get("period_return_pct"))),
        ("Period high", f"{_usd(sm.get('max_price'))} ({sm.get('max_date')})"),
        ("Period low", f"{_usd(sm.get('min_price'))} ({sm.get('min_date')})"),
        ("Obs-to-obs volatility",
         f"{sm['observation_to_observation_volatility_pct']:.2f}%"
         if sm.get("observation_to_observation_volatility_pct") is not None else "—"),
        ("Observations", str(sm.get("observations"))),
    ])

    heading("4. Market Analysis")
    para(payload["market_findings"].get("summary", ""))
    for o in payload["market_findings"].get("observations", []):
        bullet(o)

    heading("5. Fundamental Research")
    para(payload["research_findings"].get("summary", ""))
    kv([
        (f"Revenue (FY{f.get('fiscal_year') or ''})", _big(f.get("revenue"))),
        ("Net income", _big(f.get("net_income"))),
        ("Net margin", f"{fm['net_margin_pct']:.2f}%" if fm.get("net_margin_pct") else "—"),
        ("Total debt", _big(f.get("total_debt"))),
        ("Debt / revenue", f"{fm['debt_to_revenue']:.2f}x" if fm.get("debt_to_revenue") else "—"),
        ("Period end", str(f.get("period_end") or "—")),
    ])

    heading("6. Combined Analyst View (descriptive, not advice)")
    if sm.get("period_return_pct") is not None and f.get("pe_ratio"):
        para(f"Over the selected window the price trend is "
             f"{'down' if sm['period_return_pct'] < 0 else 'up'} {_pct(sm['period_return_pct'])}, "
             f"ending at {_usd(sm.get('end_price'))} within a 52-week range of "
             f"{_usd(q.get('fifty_two_week_low'))}–{_usd(q.get('fifty_two_week_high'))}. "
             f"Valuation shows a trailing P/E of {f['pe_ratio']:.1f} with an FY{f.get('fiscal_year')} "
             f"net margin of {fm.get('net_margin_pct', 0):.1f}%. These are observations from the data, "
             f"not a buy/sell recommendation or price target.")

    heading("7. Risks")
    for r in payload["risks"]:
        bullet(r)
    heading("8. Limitations")
    for l in payload["limitations"]:
        bullet(l)
    heading("9. Data Sources & Provenance")
    for s in payload["sources"]:
        bullet(s)
    para(f"Produced by a multi-agent workflow on {payload.get('model')}; "
         f"{payload['cost']['llm_call_count']} model calls, estimated "
         f"${float(payload['cost']['known_estimated_cost_usd'] or 0):.4f}. "
         f"Live Yahoo Finance data (may be ~15 min delayed), not audited.", italic=True)
    para("Not financial advice. Automated analysis for information only; not affiliated with or "
         "endorsed by the company; no buy/sell recommendation or price target.", italic=True)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
